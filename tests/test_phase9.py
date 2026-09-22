from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import tempfile

import openpyxl
import streamlit as st

from utils import auth
from utils.db import SQLiteDB
from utils.intelligence import build_risk_intelligence
from utils.performance import profile_call
from utils.assurance import validate_transition
from utils.workflow import completion_percent, stage_results


def _complete_workflow_session() -> dict:
    session = {
        "project_name": "Phase 9 Project",
        "field_name": "Phase 9 Field",
        "project_phase": "PGR2",
        "uncertainties": [
            {
                "uncertainty_id": "UNC-001",
                "discipline": "Geology",
                "name": "Fault seal",
                "selected": True,
            }
        ],
        "key_decisions": [
            {
                "Key Decision": "Well placement",
                "Weight (1-3)": 3,
            }
        ],
        "impact_assessment": [
            {
                "uncertainty_id": "UNC-001",
                "Uncertainty": "Fault seal",
                "Degree of Uncertainty": "H",
                "Well placement": "H",
                "Impact (Weighted)": 3.0,
                "Impact Bin": "H",
                "Combined Rating": "HH",
            }
        ],
        "key_uncertainties": [
            {
                "uncertainty_id": "UNC-001",
                "Uncertainty": "Fault seal",
                "Degree of Uncertainty": "H",
                "Impact (Weighted)": 3.0,
                "Combined Rating": "HH",
                "Rank": 1,
                "Include in Plan": True,
            }
        ],
        "_mapping": {
            "resolution_options": ["Fault Seal Analysis"],
        },
        "resolution_list": {
            "Fault seal": {"Fault Seal Analysis": "Y"},
        },
        "resolution_planner": [
            {
                "resolution_id": "RES-001",
                "Resolution Action": "Fault Seal Analysis",
                "Action Owner": "Engineer A",
                "Part of Workplan": True,
            }
        ],
        "risk_register": [
            {
                "risk_id": "RSK-001",
                "Risk": "Low in place volume",
                "Uncertainty/Causes": "Fault seal",
                "Resolution Plan": "Fault Seal Analysis",
                "Action Owner": "Engineer A",
                "Contingency Plan": "Review well placement",
                "Impact/Consequence": "Lower volume",
                "Likelihood (H/M/L)": "H",
                "Impact (H/M/L)": "M",
            }
        ],
        "pra_output": [{"Risk ID": "RSK-001", "Risk": "Low in place volume"}],
    }
    return session


def test_phase9_end_to_end_workflow_reaches_complete_state():
    session = _complete_workflow_session()
    stages = stage_results(session)

    assert all(stage.complete for stage in stages)
    assert completion_percent(session) == 100
    assert [stage.key for stage in stages] == [
        "uncertainties",
        "key_decisions",
        "impact_assessment",
        "key_uncertainties",
        "resolution_list",
        "resolution_planner",
        "risk_register",
        "pra_output",
    ]


def test_phase9_export_contains_all_hardened_sheets():
    from utils.export_excel import build_excel_export

    st.session_state.clear()
    for key, value in _complete_workflow_session().items():
        st.session_state[key] = value
    for key, value in {
        "study_id": "phase9-export",
        "study_owner": "Engineer A",
        "methodology_version": "SURM-2026.01",
        "study_lifecycle": "Reviewed",
        "study_revision": 3,
        "study_change_log": [{"revision": 3, "actor": "Engineer A"}],
        "team_members": [],
        "bowtie_register": {},
        "barrier_register": {},
        "study_reviews": [],
        "workflow_revisions": {},
        "workflow_snapshots": {},
    }.items():
        st.session_state[key] = value

    workbook_bytes = build_excel_export()
    workbook = openpyxl.load_workbook(BytesIO(workbook_bytes), read_only=True)

    assert workbook.sheetnames == [
        "Front Page",
        "Documentation",
        "1. Uncertainties List",
        "2. Key Decisions",
        "3. Impact Assessment",
        "4. Key Uncertainties",
        "5. Resolution List",
        "6. Resolution Planner",
        "7. Risk Register",
        "7b. Bowtie Register",
        "8. Barrier Management",
        "9. Assurance & Reviews",
        "10. Traceability",
        "PRA Output",
    ]

    assert workbook["7. Risk Register"].max_row >= 2
    assert workbook["10. Traceability"].max_row >= 2
    workbook.close()


def test_phase9_sqlite_stale_save_is_rejected():
    with tempfile.TemporaryDirectory() as directory:
        database = SQLiteDB()
        database.db_path = f"{directory}\\conflict.db"
        database.init()

        first = {
            "session": {
                "project_name": "Alpha",
                "field_name": "Beta",
                "study_revision": 1,
            },
            "meta": {
                "project_phase": "PGR1",
                "completion": 10,
                "auto_saved": False,
                "saved_at": "2026-09-22T12:00:00",
            },
        }
        second = {
            "session": {
                "project_name": "Alpha",
                "field_name": "Beta",
                "study_revision": 2,
            },
            "meta": {
                "project_phase": "PGR1",
                "completion": 20,
                "auto_saved": False,
                "saved_at": "2026-09-22T12:01:00",
            },
        }

        assert database.save_bundle("Alpha", "Beta", 1, first, expected_previous_revision=0)
        assert not database.save_bundle(
            "Alpha",
            "Beta",
            2,
            second,
            expected_previous_revision=0,
        )
        assert database.save_bundle(
            "Alpha",
            "Beta",
            2,
            second,
            expected_previous_revision=1,
        )

        assert database.load("Alpha", "Beta")["session"]["study_revision"] == 2


def test_phase9_sqlite_concurrent_writers_are_serialized():
    with tempfile.TemporaryDirectory() as directory:
        db_path = f"{directory}\\concurrent.db"

        def write(index: int) -> bool:
            database = SQLiteDB()
            database.db_path = db_path
            database.init()
            record = {
                "session": {
                    "project_name": f"Project {index}",
                    "field_name": "Shared Field",
                    "study_revision": 1,
                },
                "meta": {
                    "project_phase": "PGR1",
                    "completion": index,
                    "auto_saved": False,
                    "saved_at": f"2026-09-22T12:0{index}:00",
                },
            }
            return database.save_bundle(
                f"Project {index}",
                "Shared Field",
                1,
                record,
                expected_previous_revision=0,
            )

        with ThreadPoolExecutor(max_workers=4) as executor:
            results = list(executor.map(write, range(4)))

        database = SQLiteDB()
        database.db_path = db_path
        assert all(results)
        assert len(database.list_all_records()) == 4


def test_phase9_authenticated_identity_controls_roles(monkeypatch):
    identity = auth.Identity(
        subject="user-001",
        email="engineer@example.com",
        name="Engineer A",
        roles=frozenset({"Author", "Reviewer"}),
        authenticated=True,
        source="test",
    )

    monkeypatch.setenv("SURM_AUTH_REQUIRED", "1")

    with monkeypatch.context() as isolated:
        isolated.setattr(auth, "_streamlit_identity", lambda: identity)
        allowed, reason = auth.can_edit_study(
            {
                "study_role": "Author",
                "study_owner": "engineer@example.com",
            }
        )
        assert allowed, reason

        denied, reason = auth.can_edit_study(
            {
                "study_role": "Approver",
                "study_owner": "engineer@example.com",
            }
        )
        assert not denied
        assert "Approver" in reason


def test_phase9_unauthenticated_deployment_is_locked(monkeypatch):
    monkeypatch.setenv("SURM_AUTH_REQUIRED", "1")

    with monkeypatch.context() as isolated:
        isolated.setattr(
            auth,
            "_streamlit_identity",
            lambda: None,
        )
        allowed, reason = auth.can_edit_study(
            {
                "study_role": "Author",
                "study_owner": "local-user",
            }
        )

    assert not allowed
    assert "Authenticated identity" in reason


def test_phase9_intelligence_profiles_larger_study():
    session = {
        "risk_register": [
            {
                "risk_id": f"RSK-{i:04d}",
                "Risk": f"Risk {i}",
                "Likelihood (H/M/L)": "H",
                "Impact (H/M/L)": "M",
                "Risk Rating": "High",
                "Risk Status": "Open",
            }
            for i in range(500)
        ],
        "resolution_planner": [
            {
                "resolution_id": f"RES-{i:04d}",
                "Resolution Action": f"Action {i}",
                "Status": "In Progress",
            }
            for i in range(500)
        ],
        "bowtie_register": {},
        "barrier_register": {},
        "uncertainties": [],
    }

    result = profile_call(
        "risk-intelligence-500",
        build_risk_intelligence,
        session,
    )

    assert result.elapsed_ms >= 0
    assert result.result["risk_count"] == 500
    assert result.result["action_count"] == 500


def test_phase9_lifecycle_transition_matrix_blocks_skips_and_backwards_moves():
    cases = [
        ("Draft", "In Review", "Author", True),
        ("Draft", "Reviewed", "Reviewer", False),
        ("Reviewed", "Approved", "Reviewer", False),
        ("Reviewed", "Archived", "Approver", False),
        ("Approved", "Archived", "Approver", True),
        ("Approved", "Draft", "Approver", False),
    ]

    for current, target, role, expected in cases:
        session = _complete_workflow_session()
        session.update(
            {
                "study_lifecycle": current,
                "study_role": role,
                "study_access_mode": "edit",
            }
        )
        allowed, reasons = validate_transition(session, target)
        assert allowed is expected, (current, target, role, reasons)
